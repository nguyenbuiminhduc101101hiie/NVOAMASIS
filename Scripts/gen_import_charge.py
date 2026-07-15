import os
import openpyxl

wb = openpyxl.load_workbook(r"c:\Users\PCPV\Downloads\DANH MUC PHI.xlsx", data_only=True)
ws = wb.active


def sql_escape(v):
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    return s.replace("'", "''")


lines = []
lines.append("-- Import CHARGE from DANH MUC PHI.xlsx (from row 8)")
lines.append("-- Mapping: CHARGE_CODE=D, CHARGE=unit=C, DVT=B, APPROVE=CONTINUED=EDITABLE=1")
lines.append("SET NOCOUNT ON;")
lines.append("BEGIN TRAN;")
lines.append("")
lines.append("INSERT INTO [dbo].[CHARGE] (")
lines.append("    [CHARGE_ID], [stt], [CHARGE_CODE], [CHARGE], [unit], [DVT],")
lines.append("    [APPROVE], [CONTINUED], [EDITABLE], [UPDATETIME]")
lines.append(")")
lines.append("VALUES")

values = []
stt = 0
skipped = 0
for r in range(8, ws.max_row + 1):
    b = sql_escape(ws.cell(r, 2).value)
    c = sql_escape(ws.cell(r, 3).value)
    d = sql_escape(ws.cell(r, 4).value)
    if not d and not c:
        skipped += 1
        continue
    stt += 1
    code = f"N'{d}'" if d else "NULL"
    charge = f"N'{c}'" if c else "NULL"
    unit = f"N'{c}'" if c else "NULL"
    dvt = f"N'{b}'" if b else "NULL"
    values.append(
        f"    (NEWID(), {stt}, {code}, {charge}, {unit}, {dvt}, 1, 1, 1, GETDATE())"
    )

lines.append(",\n".join(values) + ";")
lines.append("")
lines.append(f"-- Imported {stt} rows, skipped empty: {skipped}")
lines.append("COMMIT;")
lines.append("-- ROLLBACK; -- dung de test, comment COMMIT o tren neu can rollback")

out = os.path.join(os.path.dirname(__file__), "Import_CHARGE_FromExcel.sql")
with open(out, "w", encoding="utf-8-sig") as f:
    f.write("\n".join(lines))

print(f"rows={stt} skipped={skipped}")
print(f"file={out}")
